"""Verify the built deck and print the text of every slide.

Checks, all against the saved .pptx rather than the build script's own state:
  1. exactly five slides
  2. every slide has a non-empty title placeholder
  3. every slide has non-empty speaker notes
  4. no text frame overflows its shape (word-wrapped height against box height,
     measured with the real Arial metrics PowerPoint will use)
  5. nothing extends past the slide edges
  6. every embedded image is byte-identical to a file that exists in deck/figures
  7. the live-demo URL appears on every slide, and the QR code embedded in the
     .pptx decodes back to it

Run:
    .venv/bin/python deck/build/verify_deck.py
"""

from __future__ import annotations

import hashlib
import sys
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.path.insert(0, str(Path(__file__).resolve().parent))

import qrgen  # noqa: E402  -- local module, imported after the path is set

ROOT = Path("/Users/prathmeshwalimbe/Downloads/JSW-PS1")
PPTX_PATH = ROOT / "deck" / "JSW_Surface_Defect_Detection.pptx"
FIGDIR = ROOT / "deck" / "figures"
EMU = 914400.0
SLIDE_W, SLIDE_H = 13.333, 7.5

LIVE_URL = "https://surface-vision.github.io"
LIVE_HOST = "surface-vision.github.io"

ARIAL = "/System/Library/Fonts/Supplemental/Arial.ttf"
ARIAL_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
_CACHE: dict[tuple[bool, int], ImageFont.FreeTypeFont] = {}

failures: list[str] = []
warnings: list[str] = []


def font(size_pt: float, bold: bool) -> ImageFont.FreeTypeFont:
    key = (bold, int(round(size_pt * 4)))
    if key not in _CACHE:
        _CACHE[key] = ImageFont.truetype(ARIAL_BOLD if bold else ARIAL,
                                         size=int(round(size_pt * 4)))
    return _CACHE[key]


def width_in(text: str, size_pt: float, bold: bool) -> float:
    return font(size_pt, bold).getlength(text) / 4.0 / 72.0


def frame_height_in(shape) -> float:
    """Wrapped height of a shape's text, in inches."""
    frame = shape.text_frame
    box_w = shape.width / EMU
    box_w -= (frame.margin_left + frame.margin_right) / EMU
    total = 0.0
    for paragraph in frame.paragraphs:
        tokens = []
        max_size = 0.0
        for piece in paragraph.runs:
            size = piece.font.size.pt if piece.font.size else 18.0
            bold = bool(piece.font.bold)
            max_size = max(max_size, size)
            for word in piece.text.split():
                tokens.append((word, size, bold))
        if not tokens:
            total += (max_size or 11.0) * 1.22 / 72.0
            continue
        lines, current = 1, 0.0
        space = width_in(" ", tokens[0][1], False)
        for word, size, bold in tokens:
            word_w = width_in(word, size, bold)
            gap = space if current else 0.0
            if current + gap + word_w > box_w and current:
                lines += 1
                current = word_w
            else:
                current += gap + word_w
        total += lines * max_size * 1.22 / 72.0
        if paragraph.space_after is not None:
            total += paragraph.space_after.pt / 72.0
    total += (frame.margin_top + frame.margin_bottom) / EMU
    return total


def check_images(prs: Presentation) -> list[tuple[int, str, float, float]]:
    on_disk = {}
    for path in sorted(FIGDIR.glob("*.png")):
        on_disk[hashlib.sha256(path.read_bytes()).hexdigest()] = path
    placed = []
    for index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            digest = hashlib.sha256(shape.image.blob).hexdigest()
            match = on_disk.get(digest)
            if match is None:
                failures.append(
                    f"slide {index}: embedded image matches no file in deck/figures")
                name = "<unmatched>"
            elif not match.exists():
                failures.append(f"slide {index}: {match} does not exist")
                name = match.name
            else:
                name = match.name
            placed.append((index, name, shape.width / EMU, shape.height / EMU))
    return placed


def slide_text(slide) -> list[str]:
    lines = []
    for shape in sorted(slide.shapes, key=lambda s: (round(s.top / EMU, 2), s.left)):
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(piece.text for piece in paragraph.runs).strip()
            if text:
                lines.append(text)
    return lines


def main() -> int:
    prs = Presentation(str(PPTX_PATH))
    slides = list(prs.slides)

    if len(slides) != 5:
        failures.append(f"expected 5 slides, found {len(slides)}")
    if abs(prs.slide_width / EMU - SLIDE_W) > 0.01 or abs(prs.slide_height / EMU - SLIDE_H) > 0.01:
        failures.append("slide size is not 16:9 13.333 x 7.5 in")

    for index, slide in enumerate(slides, start=1):
        title = slide.shapes.title
        if title is None or not title.text_frame.text.strip():
            failures.append(f"slide {index}: no title")
        if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
            failures.append(f"slide {index}: no speaker notes")

        for shape in slide.shapes:
            left, top = shape.left / EMU, shape.top / EMU
            right, bottom = left + shape.width / EMU, top + shape.height / EMU
            if left < -0.01 or top < -0.01 or right > SLIDE_W + 0.01 or bottom > SLIDE_H + 0.01:
                failures.append(
                    f"slide {index}: shape '{shape.name}' extends past the slide "
                    f"({left:.2f},{top:.2f})-({right:.2f},{bottom:.2f})")
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            needed = frame_height_in(shape)
            box_h = shape.height / EMU
            if needed > box_h + 0.02:
                failures.append(
                    f"slide {index}: text overflows '{shape.name}' -- needs "
                    f"{needed:.2f} in, box is {box_h:.2f} in "
                    f"[{shape.text_frame.text[:52]!r}]")

    placed = check_images(prs)

    # the QR a judge will point a phone at: decoded out of the .pptx itself
    qr_payloads: list[str] = []
    for index, slide in enumerate(slides, start=1):
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            blob = shape.image.blob
            if hashlib.sha256(blob).hexdigest() != hashlib.sha256(
                    (FIGDIR / "fig7_live_demo_qr.png").read_bytes()).hexdigest():
                continue
            try:
                payload = qrgen.decode_png(blob)
            except Exception as exc:                       # noqa: BLE001
                failures.append(f"slide {index}: embedded QR does not decode ({exc})")
                continue
            qr_payloads.append(payload)
            if payload != LIVE_URL:
                failures.append(
                    f"slide {index}: embedded QR decodes to {payload!r}, not {LIVE_URL!r}")
    if not qr_payloads:
        failures.append("no QR code is embedded in the deck")

    # the URL, in readable text, on every slide
    url_hits: dict[int, int] = {}
    for index, slide in enumerate(slides, start=1):
        text = "\n".join(slide_text(slide))
        url_hits[index] = text.count(LIVE_HOST)
        if not url_hits[index]:
            failures.append(f"slide {index}: the live-demo URL does not appear on the slide")

    print("=" * 78)
    print("DECK VERIFICATION")
    print("=" * 78)
    print(f"file        {PPTX_PATH}")
    print(f"size        {PPTX_PATH.stat().st_size / 1024:.0f} KB")
    print(f"slides      {len(slides)}")
    print(f"dimensions  {prs.slide_width / EMU:.3f} x {prs.slide_height / EMU:.3f} in (16:9)")
    print()
    print(f"live demo   {LIVE_URL}")
    print(f"            QR codes embedded: {len(qr_payloads)}, decoded payload(s): "
          f"{', '.join(sorted(set(qr_payloads))) or 'none'}")
    print("            URL mentions per slide: "
          + ", ".join(f"s{k} x{v}" for k, v in sorted(url_hits.items())))
    print()
    print("Embedded images (all byte-identical to a file in deck/figures):")
    for index, name, width, height in placed:
        path = FIGDIR / name
        print(f"  slide {index}  {name:30s} placed {width:5.2f} x {height:4.2f} in   "
              f"exists={path.exists()}  {path.stat().st_size // 1024 if path.exists() else 0} KB")
    print()

    for index, slide in enumerate(slides, start=1):
        print("-" * 78)
        print(f"SLIDE {index}  |  title: {slide.shapes.title.text_frame.text}")
        print("-" * 78)
        for text in slide_text(slide):
            print(f"    {text}")
        notes = slide.notes_slide.notes_text_frame.text.strip()
        print(f"\n    [speaker notes: {len(notes)} characters, "
              f"{len(notes.split())} words]")
        print()

    print("=" * 78)
    if warnings:
        print("WARNINGS")
        for item in warnings:
            print("  -", item)
    if failures:
        print("FAILED")
        for item in failures:
            print("  -", item)
        return 1
    print("ALL CHECKS PASS: 5 slides, every slide titled, every slide has speaker notes,")
    print("no text frame overflows its box, nothing crosses a slide edge, every embedded")
    print("image is byte-identical to an existing file in deck/figures, the live-demo URL")
    print("is on every slide, and the QR embedded in the .pptx decodes to it.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
